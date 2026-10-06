from docx import Document


def create_docx(content, filename="legal_document.docx"):
	doc = Document()

	doc.add_heading("Legal Document", level=1)

	for paragraph in content.split("\n"):
		if paragraph.strip():
			doc.add_paragraph(paragraph)

	doc.save(filename)

	return filename
